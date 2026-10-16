"""
Vehicle History router.

Endpoints:
  GET  /vehicle-history/{vehicle_id}          — история обслуживания автомобиля (JSON)
  GET  /vehicle-history/{vehicle_id}/export   — экспорт истории в красивый DOCX
"""
import os
from datetime import datetime
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel, ConfigDict
from sqlalchemy.orm import Session

from database import get_db
from access import OrderAccess, get_order_access
from routers.settings import get_company_name
import models

router = APIRouter()

GENERATED_DIR = os.path.join(os.path.dirname(__file__), "..", "generated")
os.makedirs(GENERATED_DIR, exist_ok=True)

STATUS_LABELS = {
    "new": "Новый",
    "in_progress": "В работе",
    "done": "Готов",
    "issued": "Выдан",
    "paid": "Оплачен",
}


# ─── Schemas ─────────────────────────────────────────────────────────────────

class ServiceHistoryItem(BaseModel):
    order_id: int
    date: datetime
    mileage: Optional[int]
    status: str
    works: List[str]
    parts: List[str]
    services_total: float
    parts_total: float
    total_amount: float
    comment: Optional[str]


class VehicleHistoryOut(BaseModel):
    vehicle_id: int
    brand: str
    model: str
    plate_number: Optional[str]
    vin: Optional[str]
    year: Optional[int]
    client_name: Optional[str]
    client_phone: Optional[str]
    orders: List[ServiceHistoryItem]
    total_visits: int
    total_spent: float


# ─── API Endpoints ────────────────────────────────────────────────────────────

@router.get("/{vehicle_id}", response_model=VehicleHistoryOut)
def get_vehicle_history(
    vehicle_id: int,
    db: Session = Depends(get_db),
    access: OrderAccess = Depends(get_order_access),
):
    # История содержит все работы и суммы по авто — роли «только свои работы» она недоступна
    access.require_full_access()
    """Возвращает историю обслуживания автомобиля."""
    vehicle = db.query(models.Vehicle).filter(models.Vehicle.id == vehicle_id).first()
    if not vehicle:
        raise HTTPException(status_code=404, detail="Автомобиль не найден")

    orders = (
        db.query(models.Order)
        .filter(models.Order.vehicle_id == vehicle_id)
        .order_by(models.Order.created_at.desc())
        .all()
    )

    history_items = []
    total_spent = 0.0

    for order in orders:
        services_total = round(sum(item.price for item in order.items), 2)
        parts_total = round(sum(p.price * p.qty for p in order.parts), 2)
        total = round(services_total + parts_total, 2)
        total_spent += total

        works = [item.work_name for item in order.items]
        parts = [part.name for part in order.parts]

        history_items.append(ServiceHistoryItem(
            order_id=order.id,
            date=order.created_at,
            mileage=order.mileage,
            status=STATUS_LABELS.get(order.status.value, order.status.value),
            works=works,
            parts=parts,
            services_total=services_total,
            parts_total=parts_total,
            total_amount=total,
            comment=order.comment,
        ))

    return VehicleHistoryOut(
        vehicle_id=vehicle.id,
        brand=vehicle.brand,
        model=vehicle.model,
        plate_number=vehicle.plate_number,
        vin=vehicle.vin,
        year=vehicle.year,
        client_name=vehicle.client.full_name if vehicle.client else None,
        client_phone=vehicle.client.phone if vehicle.client else None,
        orders=history_items,
        total_visits=len(history_items),
        total_spent=round(total_spent, 2),
    )


@router.get("/{vehicle_id}/export")
def export_vehicle_history(
    vehicle_id: int,
    db: Session = Depends(get_db),
    access: OrderAccess = Depends(get_order_access),
):
    # История содержит все работы и суммы по авто — роли «только свои работы» она недоступна
    access.require_full_access()
    """Генерирует красивый DOCX-отчёт истории обслуживания автомобиля."""
    vehicle = db.query(models.Vehicle).filter(models.Vehicle.id == vehicle_id).first()
    if not vehicle:
        raise HTTPException(status_code=404, detail="Автомобиль не найден")

    orders = (
        db.query(models.Order)
        .filter(models.Order.vehicle_id == vehicle_id)
        .order_by(models.Order.created_at.asc())
        .all()
    )

    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
    except ImportError:
        raise HTTPException(status_code=500, detail="Библиотека python-docx не установлена")

    doc = Document()

    # ─── Page margins ─────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2)

    def set_font(run, size=10, bold=False, color=None):
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor(*color)

    def add_styled_para(text, size=10, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        set_font(run, size, bold, color)
        return p

    def set_cell_background(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}" w:color="auto" w:val="clear"/>')
        tcPr.append(shd)

    def set_cell_padding(cell, top=60, bottom=60, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'  <w:top w:w="{top}" w:type="dxa"/>'
            f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'  <w:left w:w="{left}" w:type="dxa"/>'
            f'  <w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    def set_cell_borders(cell, color="DDDDDD"):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:color="{color}" w:sz="4"/>'
            f'  <w:left w:val="single" w:color="{color}" w:sz="4"/>'
            f'  <w:bottom w:val="single" w:color="{color}" w:sz="4"/>'
            f'  <w:right w:val="single" w:color="{color}" w:sz="4"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

    # ─── Header ───────────────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_run = title_p.add_run("ИСТОРИЯ ОБСЛУЖИВАНИЯ")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(20)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(14)
    sub_run = subtitle_p.add_run(
        f"{get_company_name(db)} · Отчёт сформирован: " + datetime.now().strftime("%d.%m.%Y %H:%M")
    )
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(9)
    sub_run.font.color.rgb = RGBColor(0x88, 0x88, 0x99)

    # ─── Vehicle Info Card ────────────────────────────────────────────────────
    add_styled_para("АВТОМОБИЛЬ", size=8, bold=True, color=(0x88, 0x88, 0x99), space_after=4)

    info_table = doc.add_table(rows=0, cols=4)
    info_table.style = "Table Grid"

    client_name = vehicle.client.full_name if vehicle.client else "—"
    client_phone = vehicle.client.phone if vehicle.client else "—"
    plate = vehicle.plate_number or "—"
    vin = vehicle.vin or "—"
    year = str(vehicle.year) if vehicle.year else "—"
    vehicle_label = f"{vehicle.brand} {vehicle.model}"

    info_rows = [
        ("Марка / Модель:", vehicle_label, "Гос. номер:", plate),
        ("VIN:", vin, "Год выпуска:", year),
        ("Клиент:", client_name, "Телефон:", client_phone),
    ]

    for label1, val1, label2, val2 in info_rows:
        row = info_table.add_row()
        for cell, text, is_label in [
            (row.cells[0], label1, True),
            (row.cells[1], val1, False),
            (row.cells[2], label2, True),
            (row.cells[3], val2, False),
        ]:
            cell.text = text
            if is_label:
                set_cell_background(cell, "F5F5F7")
            set_cell_padding(cell)
            set_cell_borders(cell)
            for p in cell.paragraphs:
                for r in p.runs:
                    if is_label:
                        set_font(r, 9, bold=True, color=(0x44, 0x44, 0x55))
                    else:
                        set_font(r, 9)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ─── Summary stats ────────────────────────────────────────────────────────
    total_spent = sum(
        round(sum(i.price for i in o.items) + sum(p.price * p.qty for p in o.parts), 2)
        for o in orders
    )
    add_styled_para("СТАТИСТИКА", size=8, bold=True, color=(0x88, 0x88, 0x99), space_after=4)

    stats_table = doc.add_table(rows=1, cols=3)
    stats_cells = stats_table.rows[0].cells
    stats_data = [
        ("Всего визитов", str(len(orders))),
        ("Потрачено всего", f"{total_spent:,.0f} ₸".replace(",", " ")),
        ("Последнее ТО", orders[-1].created_at.strftime("%d.%m.%Y") if orders else "—"),
    ]
    for cell, (label, val) in zip(stats_cells, stats_data):
        p1 = cell.paragraphs[0]
        p1.clear()
        r1 = p1.add_run(val)
        set_font(r1, 18, bold=True, color=(0x1A, 0x1A, 0x2E))
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2 = cell.add_paragraph(label)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p2.runs:
            set_font(r, 8, color=(0x88, 0x88, 0x99))
        set_cell_background(cell, "F0F0F5")
        set_cell_padding(cell, top=120, bottom=120, left=80, right=80)
        set_cell_borders(cell, "DDDDEE")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ─── Orders history ───────────────────────────────────────────────────────
    add_styled_para("ИСТОРИЯ ВИЗИТОВ", size=8, bold=True, color=(0x88, 0x88, 0x99), space_after=6)

    if not orders:
        add_styled_para("Заказ-нарядов по этому автомобилю не найдено.", size=10, color=(0x88, 0x88, 0x99))
    else:
        for order in reversed(orders):  # newest first
            services_total = round(sum(item.price for item in order.items), 2)
            parts_total = round(sum(p.price * p.qty for p in order.parts), 2)
            total = round(services_total + parts_total, 2)

            # ── Order header ──
            header_table = doc.add_table(rows=1, cols=3)
            hrow = header_table.rows[0]

            cell_num = hrow.cells[0]
            cell_num.text = ""
            p = cell_num.paragraphs[0]
            r_num = p.add_run(f"Заказ-наряд #{order.id}")
            set_font(r_num, 11, bold=True, color=(0xFF, 0xFF, 0xFF))
            set_cell_background(cell_num, "1A1A2E")
            set_cell_padding(cell_num, top=80, bottom=80, left=120, right=60)

            cell_date = hrow.cells[1]
            cell_date.text = ""
            p2 = cell_date.paragraphs[0]
            date_str = order.created_at.strftime("%d.%m.%Y")
            mileage_str = (
                f"  ·  Пробег: {order.mileage:,} км".replace(",", " ")
                if order.mileage else ""
            )
            r_date = p2.add_run(f"{date_str}{mileage_str}")
            set_font(r_date, 9, color=(0xCC, 0xCC, 0xDD))
            set_cell_background(cell_date, "1A1A2E")
            set_cell_padding(cell_date, top=80, bottom=80, left=60, right=60)

            cell_total = hrow.cells[2]
            cell_total.text = ""
            p3 = cell_total.paragraphs[0]
            p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_total = p3.add_run(f"{total:,.0f} ₸".replace(",", " "))
            set_font(r_total, 12, bold=True, color=(0xF0, 0xD0, 0x60))
            set_cell_background(cell_total, "1A1A2E")
            set_cell_padding(cell_total, top=80, bottom=80, left=60, right=120)

            # ── Works table ──
            if order.items:
                works_table = doc.add_table(rows=1, cols=3)
                hdr = works_table.rows[0].cells
                for i, (cell, text) in enumerate(zip(hdr, ["Наименование работы", "Исполнитель", "Сумма (₸)"])):
                    cell.text = text
                    set_cell_background(cell, "EEEEF5")
                    set_cell_padding(cell, top=60, bottom=60, left=100, right=60)
                    set_cell_borders(cell, "CCCCDD")
                    for p in cell.paragraphs:
                        for r in p.runs:
                            set_font(r, 8, bold=True, color=(0x33, 0x33, 0x55))
                    if i == 2:
                        for p in cell.paragraphs:
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

                for item in order.items:
                    row = works_table.add_row()
                    row.cells[0].text = item.work_name
                    row.cells[1].text = item.employee.full_name if item.employee else "—"
                    row.cells[2].text = f"{item.price:,.0f}".replace(",", " ")
                    for j, cell in enumerate(row.cells):
                        set_cell_padding(cell, top=50, bottom=50, left=100, right=60)
                        set_cell_borders(cell, "DDDDEE")
                        for p in cell.paragraphs:
                            for r in p.runs:
                                set_font(r, 9)
                        if j == 2:
                            for p in cell.paragraphs:
                                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

            # ── Parts table ──
            if order.parts:
                parts_table = doc.add_table(rows=1, cols=3)
                hdr = parts_table.rows[0].cells
                for i, (cell, text) in enumerate(zip(hdr, ["Запчасть / Материал", "Кол-во", "Сумма (₸)"])):
                    cell.text = text
                    set_cell_background(cell, "FFF8EE")
                    set_cell_padding(cell, top=60, bottom=60, left=100, right=60)
                    set_cell_borders(cell, "EEDDCC")
                    for p in cell.paragraphs:
                        for r in p.runs:
                            set_font(r, 8, bold=True, color=(0x55, 0x44, 0x22))
                    if i == 2:
                        for p in cell.paragraphs:
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

                for part in order.parts:
                    row = parts_table.add_row()
                    row.cells[0].text = part.name
                    row.cells[1].text = str(part.qty).rstrip("0").rstrip(".")
                    row.cells[2].text = f"{part.price * part.qty:,.0f}".replace(",", " ")
                    for j, cell in enumerate(row.cells):
                        set_cell_padding(cell, top=50, bottom=50, left=100, right=60)
                        set_cell_borders(cell, "EEDDCC")
                        for p in cell.paragraphs:
                            for r in p.runs:
                                set_font(r, 9)
                        if j == 2:
                            for p in cell.paragraphs:
                                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

            # ── Comment ──
            if order.comment:
                comment_p = doc.add_paragraph()
                comment_p.paragraph_format.space_before = Pt(3)
                comment_p.paragraph_format.space_after = Pt(3)
                r_lbl = comment_p.add_run("Комментарий: ")
                set_font(r_lbl, 8, bold=True, color=(0x77, 0x77, 0x88))
                r_cmt = comment_p.add_run(order.comment)
                set_font(r_cmt, 8, color=(0x44, 0x44, 0x55))

            # Gap between orders
            gap = doc.add_paragraph()
            gap.paragraph_format.space_before = Pt(0)
            gap.paragraph_format.space_after = Pt(8)

    # ─── Footer ───────────────────────────────────────────────────────────────
    footer_p = doc.add_paragraph()
    footer_p.paragraph_format.space_before = Pt(12)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run(
        f"Документ сформирован автоматически · {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    )
    set_font(footer_run, 7, color=(0xBB, 0xBB, 0xCC))

    # ─── Save & return ────────────────────────────────────────────────────────
    safe_plate = (vehicle.plate_number or str(vehicle.id)).replace(" ", "_").replace("/", "-")
    out_filename = f"history_{safe_plate}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    out_path = os.path.join(GENERATED_DIR, out_filename)
    doc.save(out_path)

    display_name = f"История_{vehicle.brand}_{vehicle.model}_{safe_plate}.docx"

    return FileResponse(
        out_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=display_name,
    )
